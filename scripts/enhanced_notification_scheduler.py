"""
Enhanced Notification Scheduler with Auto-Removal and Excel Table Support
Removes sent notifications and maintains proper Excel table format for Power Automate
"""

import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime, timedelta
import json
import os
import time
import logging
from pathlib import Path

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('scheduler.log'),
        logging.StreamHandler()
    ]
)

class EnhancedNotificationScheduler:
    def __init__(self):
        self.network_folder = Path("G:/Test")
        self.queue_file = Path("network_folder_simplified/sent_noti_now.xlsx")
        self.network_queue_file = Path("G:/Test/sent_noti_now.xlsx")  # Network location for Power Automate
        self.tracking_file = Path("schedule_tracking.json")
        
        # Ensure directories exist
        self.queue_file.parent.mkdir(exist_ok=True, parents=True)
        self.network_folder.mkdir(exist_ok=True, parents=True)
        
        # Initialize or load tracking
        self.load_tracking()
        
        # Ensure queue files exist on initialization
        self.initialize_queue_files()
        
        # Notification configurations with exact requirements
        self.notification_configs = {
            "Daily Status Reminders": {
                "file": "01_daily_status_reminders.xlsx",
                "schedule": {"times": ["09:00", "12:00", "15:00", "18:00"], "days": "daily"},
                "priority": "Medium",
                "recipients": "Vendors",
                "message_template": "Good morning! Please submit your timesheet for today. This is reminder {count} of 4 for today.",
                "max_daily": 4
            },
            "Manager Summary": {
                "file": "02_manager_summary_notifications.xlsx",
                "schedule": {"times": ["12:00", "14:00"], "days": "daily"},
                "priority": "Medium",
                "recipients": "Managers",
                "message_template": "Manager Summary: {count} pending submissions require your review.",
                "max_daily": 2
            },
            "All-Complete": {
                "file": "03_manager_all_complete_notifications.xlsx",
                "schedule": {"type": "event", "trigger": "all_submitted"},
                "priority": "Low",
                "recipients": "Managers",
                "message_template": "Great news! All team members have submitted their timesheets.",
                "max_daily": 1
            },
            "Mismatch Alerts": {
                "file": "04_mismatch_notifications.xlsx",
                "schedule": {"times": ["18:00"], "days": "daily"},
                "priority": "High",
                "recipients": "Vendors,Managers",
                "message_template": "URGENT: {hours} hour mismatch detected in your timesheet. Please correct immediately.",
                "max_daily": 1
            },
            "Manager Feedback": {
                "file": "05_manager_feedback_notifications.xlsx",
                "schedule": {"type": "event", "trigger": "feedback_provided"},
                "priority": "Medium",
                "recipients": "Vendors",
                "message_template": "Manager has reviewed your timesheet. Please check feedback and resubmit if needed.",
                "max_daily": 5
            },
            "Monthly Reports": {
                "file": "06_monthly_report_notifications.xlsx",
                "schedule": {"times": ["09:00"], "days": "monthly", "day_of_month": 1},
                "priority": "Medium",
                "recipients": "Managers,Admins",
                "message_template": "Monthly Report Ready: Review last month's timesheet summary and analytics.",
                "max_daily": 1
            },
            "System Alerts": {
                "file": "07_admin_system_alerts.xlsx",
                "schedule": {"type": "event", "trigger": "system_issue"},
                "priority": "Critical",
                "recipients": "Admins",
                "message_template": "CRITICAL: System maintenance required. {issue_description}",
                "max_daily": 10
            },
            "Holiday Reminders": {
                "file": "08_holiday_reminder_notifications.xlsx",
                "schedule": {"type": "holiday", "days_before": [3, 1]},
                "priority": "Low",
                "recipients": "All Users",
                "message_template": "Holiday Notice: Office will be closed in {days} days. Please submit timesheets early.",
                "max_daily": 2
            },
            "Late Submission": {
                "file": "09_late_submission_alerts.xlsx",
                "schedule": {"type": "deadline", "hours_after": [24, 48]},
                "priority": "High",
                "recipients": "Managers,Admins",
                "message_template": "OVERDUE {hours}h: Employee {employee_id} has not submitted timesheet.",
                "max_daily": 2
            }
        }
    
    def load_tracking(self):
        """Load or initialize tracking data"""
        if self.tracking_file.exists():
            with open(self.tracking_file, 'r') as f:
                self.tracking = json.load(f)
        else:
            self.tracking = {}
    
    def initialize_queue_files(self):
        """Initialize queue files if they don't exist"""
        try:
            # Check if network queue file exists
            if not self.network_queue_file.exists():
                logging.info("Network queue file not found, creating...")
                self.create_empty_queue()
            
            # Check if local queue file exists
            if not self.queue_file.exists():
                logging.info("Local queue file not found, creating...")
                self.create_empty_queue()
                
        except Exception as e:
            logging.error(f"Error initializing queue files: {e}")
    
    def save_tracking(self):
        """Save tracking data"""
        with open(self.tracking_file, 'w') as f:
            json.dump(self.tracking, f, indent=2)
    
    def create_excel_table(self, file_path, df):
        """Create a properly formatted Excel file with table for Power Automate"""
        # Write DataFrame to Excel
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Notifications', index=False)
            
            # Get the workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Notifications']
            
            # Only create table if there's data
            if len(df) > 0:
                # Define the table range
                min_col = 1
                min_row = 1
                max_col = len(df.columns)
                max_row = len(df) + 1  # +1 for header
                
                # Create table reference
                table_ref = f"A1:{openpyxl.utils.get_column_letter(max_col)}{max_row}"
                
                # Create a table
                table = Table(displayName="NotificationQueue", ref=table_ref)
                
                # Add a style to the table
                style = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False
                )
                table.tableStyleInfo = style
                
                # Add the table to the worksheet
                worksheet.add_table(table)
            
            # Auto-adjust column widths
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
        
        logging.info(f"Created Excel table with {len(df)} rows at {file_path}")
        
        # Also copy to network location if this is the queue file
        if file_path == self.queue_file:
            self.sync_to_network_folder()
    
    def remove_sent_notifications(self):
        """Remove notifications where NTStatusSent is True"""
        if not self.queue_file.exists():
            logging.info("No queue file exists yet")
            return 0
        
        try:
            # Read current queue
            df = pd.read_excel(self.queue_file)
            initial_count = len(df)
            
            # Filter out sent notifications
            if 'NTStatusSent' in df.columns:
                # Keep only unsent notifications
                df_unsent = df[df['NTStatusSent'] != True].copy()
                removed_count = initial_count - len(df_unsent)
                
                if removed_count > 0:
                    logging.info(f"Removing {removed_count} sent notifications from queue")
                    
                    # Save updated queue with Excel table format
                    self.create_excel_table(self.queue_file, df_unsent)
                    
                    logging.info(f"✅ Removed {removed_count} sent notifications. {len(df_unsent)} remain in queue.")
                    return removed_count
                else:
                    logging.info("No sent notifications to remove")
                    return 0
            else:
                logging.warning("NTStatusSent column not found in queue file")
                return 0
                
        except Exception as e:
            logging.error(f"Error removing sent notifications: {e}")
            return 0
    
    def should_send_notification(self, notif_type, current_time):
        """Check if notification should be sent based on schedule"""
        config = self.notification_configs[notif_type]
        schedule = config["schedule"]
        today_key = current_time.strftime("%Y-%m-%d")
        
        # Initialize tracking for today if not exists
        if today_key not in self.tracking:
            self.tracking[today_key] = {}
        
        if notif_type not in self.tracking[today_key]:
            self.tracking[today_key][notif_type] = {
                "count": 0,
                "last_sent": None,
                "sent_times": []
            }
        
        daily_data = self.tracking[today_key][notif_type]
        max_daily = config.get("max_daily", 999)
        
        # Check if max daily limit reached
        if daily_data["count"] >= max_daily:
            return False
        
        # Time-based schedules
        if "times" in schedule:
            current_time_str = current_time.strftime("%H:%M")
            
            for scheduled_time in schedule["times"]:
                # Check if it's time and hasn't been sent for this time slot
                if current_time_str == scheduled_time:
                    time_slot_key = f"{today_key}_{scheduled_time}"
                    if time_slot_key not in daily_data.get("sent_times", []):
                        return True
        
        # Event-driven schedules
        elif schedule.get("type") == "event":
            # Check for trigger conditions (simplified for demo)
            # In production, this would check actual trigger conditions
            return self.check_event_trigger(notif_type, schedule.get("trigger"))
        
        # Monthly schedules
        elif schedule.get("days") == "monthly":
            if current_time.day == schedule.get("day_of_month", 1):
                scheduled_time = schedule["times"][0]
                if current_time.strftime("%H:%M") == scheduled_time:
                    if daily_data["count"] == 0:
                        return True
        
        # Holiday schedules
        elif schedule.get("type") == "holiday":
            # Check for upcoming holidays (simplified)
            return self.check_holiday_schedule(current_time, schedule.get("days_before", []))
        
        # Deadline schedules
        elif schedule.get("type") == "deadline":
            # Check for missed deadlines (simplified)
            return self.check_deadline_schedule(current_time, schedule.get("hours_after", []))
        
        return False
    
    def check_event_trigger(self, notif_type, trigger):
        """Check if event trigger condition is met"""
        # Simplified event checking - in production would check actual conditions
        return False  # Only trigger when actual event occurs
    
    def check_holiday_schedule(self, current_time, days_before):
        """Check holiday reminder schedule"""
        # Simplified - in production would check actual holiday calendar
        return False
    
    def check_deadline_schedule(self, current_time, hours_after):
        """Check deadline miss schedule"""
        # Simplified - in production would check actual submission deadlines
        return False
    
    def sync_to_network_folder(self):
        """Sync the queue file to network folder for Power Automate"""
        try:
            import shutil
            
            # If local queue file exists, copy it
            if self.queue_file.exists():
                shutil.copy2(self.queue_file, self.network_queue_file)
                logging.info(f"📤 Synced queue to network: {self.network_queue_file}")
                return True
            else:
                # If local file doesn't exist, create an empty one in both locations
                logging.info("Local queue file doesn't exist, creating new empty queue")
                self.create_empty_queue()
                return True
        except Exception as e:
            logging.error(f"Error syncing to network folder: {e}")
            # Try to create empty queue in network if sync fails
            try:
                self.ensure_network_queue_exists()
            except:
                pass
            return False
    
    def create_empty_queue(self):
        """Create an empty queue file with proper structure"""
        try:
            # Create empty DataFrame with required columns
            empty_df = pd.DataFrame(columns=[
                'EmployeeID',
                'ContactEmail', 
                'Message',
                'NotificationType',
                'Priority',
                'NTStatusSent'
            ])
            
            # Save to both local and network locations
            self.create_excel_table(self.queue_file, empty_df)
            
            logging.info("✅ Created empty queue files in local and network folders")
            return True
        except Exception as e:
            logging.error(f"Error creating empty queue: {e}")
            return False
    
    def ensure_network_queue_exists(self):
        """Ensure the network queue file exists, create if not"""
        try:
            if not self.network_queue_file.exists():
                # Create empty DataFrame with required columns
                empty_df = pd.DataFrame(columns=[
                    'EmployeeID',
                    'ContactEmail', 
                    'Message',
                    'NotificationType',
                    'Priority',
                    'NTStatusSent'
                ])
                
                # Save directly to network location with Excel table format
                with pd.ExcelWriter(self.network_queue_file, engine='openpyxl') as writer:
                    empty_df.to_excel(writer, sheet_name='Notifications', index=False)
                    
                    workbook = writer.book
                    worksheet = writer.sheets['Notifications']
                    
                    # Create empty table structure
                    table_ref = "A1:F1"  # Just headers for empty table
                    table = Table(displayName="NotificationQueue", ref=table_ref)
                    style = TableStyleInfo(
                        name="TableStyleMedium2",
                        showFirstColumn=False,
                        showLastColumn=False,
                        showRowStripes=True,
                        showColumnStripes=False
                    )
                    table.tableStyleInfo = style
                    worksheet.add_table(table)
                
                logging.info(f"✅ Created network queue file at: {self.network_queue_file}")
            return True
        except Exception as e:
            logging.error(f"Error ensuring network queue exists: {e}")
            return False
    
    def read_notification_file(self, file_path):
        """Read Excel file and return DataFrame"""
        try:
            if file_path.exists():
                return pd.read_excel(file_path)
            else:
                logging.warning(f"File not found: {file_path}")
                return pd.DataFrame()
        except Exception as e:
            logging.error(f"Error reading {file_path}: {e}")
            return pd.DataFrame()
    
    def add_notifications_to_queue(self, notifications_to_add):
        """Add new notifications to queue file with Excel table format"""
        # Read existing queue or create new
        if self.queue_file.exists():
            try:
                existing_df = pd.read_excel(self.queue_file)
            except:
                existing_df = pd.DataFrame()
        else:
            existing_df = pd.DataFrame()
        
        # Combine with new notifications
        if len(notifications_to_add) > 0:
            new_df = pd.DataFrame(notifications_to_add)
            combined_df = pd.concat([existing_df, new_df], ignore_index=True)
        else:
            combined_df = existing_df
        
        # Ensure NTStatusSent column exists and is False for new items
        if 'NTStatusSent' not in combined_df.columns:
            combined_df['NTStatusSent'] = False
        
        # Fill NaN values in NTStatusSent with False
        combined_df['NTStatusSent'] = combined_df['NTStatusSent'].fillna(False)
        
        # Save with Excel table format
        self.create_excel_table(self.queue_file, combined_df)
        
        # Sync to network folder
        self.sync_to_network_folder()
        
        return len(notifications_to_add)
    
    def process_notifications(self):
        """Main processing function"""
        current_time = datetime.now()
        logging.info(f"\n{'='*60}")
        logging.info(f"Processing notifications at {current_time.strftime('%Y-%m-%d %H:%M:%S')}")
        
        # First, ensure network queue file exists
        self.ensure_network_queue_exists()
        
        # First, remove any sent notifications
        removed = self.remove_sent_notifications()
        if removed > 0:
            logging.info(f"Cleaned up {removed} sent notifications")
        
        notifications_to_add = []
        
        for notif_type, config in self.notification_configs.items():
            if self.should_send_notification(notif_type, current_time):
                # Read the source file
                file_path = self.network_folder / config["file"]
                df = self.read_notification_file(file_path)
                
                if not df.empty:
                    # Process each row
                    for _, row in df.iterrows():
                        notification = {
                            'EmployeeID': row.get('EmployeeID', ''),
                            'ContactEmail': row.get('ContactEmail', ''),
                            'Message': config["message_template"].format(
                                count=self.tracking.get(current_time.strftime("%Y-%m-%d"), {})
                                     .get(notif_type, {}).get("count", 0) + 1,
                                hours=8,  # Placeholder
                                days=3,   # Placeholder
                                employee_id=row.get('EmployeeID', ''),
                                issue_description="System maintenance required"
                            ),
                            'NotificationType': notif_type,
                            'Priority': config["priority"],
                            'NTStatusSent': False
                        }
                        notifications_to_add.append(notification)
                    
                    # Update tracking
                    today_key = current_time.strftime("%Y-%m-%d")
                    current_time_str = current_time.strftime("%H:%M")
                    
                    self.tracking[today_key][notif_type]["count"] += len(df)
                    self.tracking[today_key][notif_type]["last_sent"] = current_time.isoformat()
                    
                    if "sent_times" not in self.tracking[today_key][notif_type]:
                        self.tracking[today_key][notif_type]["sent_times"] = []
                    
                    time_slot_key = f"{today_key}_{current_time_str}"
                    self.tracking[today_key][notif_type]["sent_times"].append(time_slot_key)
                    
                    logging.info(f"✅ Queued {len(df)} notifications for {notif_type}")
        
        # Add all notifications to queue
        if notifications_to_add:
            added_count = self.add_notifications_to_queue(notifications_to_add)
            logging.info(f"📧 Total notifications added to queue: {added_count}")
        else:
            logging.info("No notifications to queue at this time")
        
        # Save tracking
        self.save_tracking()
        
        # Display queue status
        self.display_queue_status()
    
    def display_queue_status(self):
        """Display current queue status"""
        if self.queue_file.exists():
            try:
                df = pd.read_excel(self.queue_file)
                total = len(df)
                
                if total > 0:
                    unsent = len(df[df['NTStatusSent'] != True]) if 'NTStatusSent' in df.columns else total
                    sent = total - unsent
                    
                    logging.info(f"\n📊 Queue Status:")
                    logging.info(f"   Total in queue: {total}")
                    logging.info(f"   Unsent: {unsent}")
                    logging.info(f"   Sent (awaiting removal): {sent}")
                    
                    # Show breakdown by type
                    if 'NotificationType' in df.columns:
                        type_counts = df[df['NTStatusSent'] != True]['NotificationType'].value_counts()
                        if not type_counts.empty:
                            logging.info(f"\n   By Type (unsent only):")
                            for notif_type, count in type_counts.items():
                                logging.info(f"   - {notif_type}: {count}")
                else:
                    logging.info("\n📊 Queue is empty")
            except Exception as e:
                logging.error(f"Error reading queue status: {e}")
    
    def run_continuous(self, interval_minutes=10):
        """Run scheduler continuously"""
        logging.info(f"🚀 Starting Enhanced Notification Scheduler")
        logging.info(f"   - Checking every {interval_minutes} minutes")
        logging.info(f"   - Auto-removing sent notifications")
        logging.info(f"   - Creating Excel tables for Power Automate")
        logging.info(f"   - Local queue: {self.queue_file}")
        logging.info(f"   - Network sync: {self.network_queue_file}")
        
        while True:
            try:
                self.process_notifications()
                
                # Show next check time
                next_check = datetime.now() + timedelta(minutes=interval_minutes)
                logging.info(f"\n⏰ Next check at {next_check.strftime('%H:%M:%S')}")
                logging.info("=" * 60)
                
                # Wait for next interval
                time.sleep(interval_minutes * 60)
                
            except KeyboardInterrupt:
                logging.info("\n🛑 Scheduler stopped by user")
                break
            except Exception as e:
                logging.error(f"Error in scheduler: {e}")
                time.sleep(60)  # Wait a minute before retrying

def main():
    scheduler = EnhancedNotificationScheduler()
    
    # Check for test mode
    import sys
    if len(sys.argv) > 1 and sys.argv[1] == '--test':
        # Run once for testing
        scheduler.process_notifications()
    else:
        # Run continuously
        scheduler.run_continuous(interval_minutes=10)

if __name__ == "__main__":
    main()
